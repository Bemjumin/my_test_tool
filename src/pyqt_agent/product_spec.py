"""产品需求说明书：按公开文章的章节写，作为确认、开发和测试的唯一原稿。"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from pyqt_agent.docx_parser import _clean_item, _compact, _heading_level, _iter_blocks, _pad, _split_label, _table_rows
from pyqt_agent.models import PLACEHOLDER, is_blank

APPENDIX_NOTE = "本附录仅供参考，不作为正式需求，不作为验收标准。"

_SECTIONS = ("首页", "修订记录", "引言", "项目概述", "业务需求", "附录", "签字")
_SECTION_ALIASES = {"修订页": "修订记录", "封面": "首页", "签字页": "签字"}
_SUBSECTIONS = {
    "引言": ("编写目的", "范围", "定义", "参考资料"),
    "项目概述": ("项目背景", "项目周期", "产品描述", "运行方式", "使用者", "产品功能"),
    "业务需求": ("不可接受的现象", "业务规则"),
}
_FEATURE_LABELS = {
    "所属模块": "module",
    "功能介绍": "intro",
    "入口": "entry",
    "前置状态": "precondition",
    "边界时会看到": "boundary_sees",
    "不可接受": "unacceptable",
    "本次不做": "excluded",
}
_COVER_FIELDS = (
    ("公司名称", "company"),
    ("文档标题", "title"),
    ("文档编号", "doc_id"),
    ("编写人", "author"),
    ("模块名称", "module"),
    ("部门", "department"),
    ("保密等级", "confidentiality"),
    ("日期", "date"),
    ("版权说明", "copyright"),
)
_SCOPE_FIELDS = (
    ("软件名称", "software_name"),
    ("软件要做的事", "does"),
    ("软件不做的事", "does_not"),
    ("应用目标", "goal"),
)
_RUNTIME_FIELDS = (
    ("exe 文件名", "exe_name"),
    ("主窗口标题", "window_title"),
    ("如何启动", "how_to_start"),
    ("启动参数", "launch_args"),
    ("如何退出", "how_to_exit"),
    ("工作目录", "workdir"),
    ("是否需要登录", "needs_login"),
    ("测试账号", "account"),
    ("分辨率", "resolution"),
)
_ITEM_PREFIX = re.compile(r"^(\d+[.、)]\s*|[-•]\s*)")
_HEADER_FILL = "1F4E79"


@dataclass
class CoverPage:
    company: str = ""
    title: str = ""
    doc_id: str = ""
    author: str = ""
    module: str = ""
    department: str = ""
    confidentiality: str = ""
    date: str = ""
    copyright: str = ""


@dataclass
class Revision:
    number: str
    chapter: str
    summary: str
    date: str
    version_before: str
    version_after: str
    author: str
    approver: str


@dataclass
class Scope:
    software_name: str = ""
    does: str = ""
    does_not: str = ""
    goal: str = ""


@dataclass
class Definition:
    term: str
    meaning: str
    ui_name: str = ""


@dataclass
class Reference:
    name: str
    note: str = ""


@dataclass
class Runtime:
    exe_name: str = ""
    window_title: str = ""
    how_to_start: str = ""
    launch_args: str = ""
    how_to_exit: str = ""
    workdir: str = ""
    needs_login: str = ""
    account: str = ""
    resolution: str = ""


@dataclass
class SpecFile:
    role: str
    path: str
    description: str


@dataclass
class UserRole:
    name: str
    who: str
    task: str
    cares: str
    avoid: str


@dataclass
class Module:
    module_id: str
    name: str
    summary: str
    priority: str


@dataclass
class FlowStep:
    applicability: str
    text: str


@dataclass
class DataItem:
    label: str
    kind: str
    source: str
    note: str
    required: str
    valid_example: str
    invalid_example: str
    boundary: str


@dataclass
class Screen:
    path: str
    sees: str
    applicability: str


@dataclass
class BusinessRule:
    description: str
    sample_input: str
    sample_output: str


@dataclass
class Feature:
    feature_id: str
    name: str
    module: str = ""
    intro: str = ""
    entry: str = ""
    precondition: str = ""
    steps: list[FlowStep] = field(default_factory=list)
    items: list[DataItem] = field(default_factory=list)
    screens: list[Screen] = field(default_factory=list)
    boundary_sees: str = ""
    unacceptable: str = ""
    excluded: str = ""
    rules: list[BusinessRule] = field(default_factory=list)


@dataclass
class Signature:
    role: str
    name: str
    date: str


@dataclass
class ProductSpec:
    cover: CoverPage = field(default_factory=CoverPage)
    revisions: list[Revision] = field(default_factory=list)
    purpose: str = ""
    scope: Scope = field(default_factory=Scope)
    definitions: list[Definition] = field(default_factory=list)
    references: list[Reference] = field(default_factory=list)
    background: str = ""
    schedule: str = ""
    description: str = ""
    runtime: Runtime = field(default_factory=Runtime)
    files: list[SpecFile] = field(default_factory=list)
    users: list[UserRole] = field(default_factory=list)
    modules: list[Module] = field(default_factory=list)
    obvious: list[str] = field(default_factory=list)
    rules: list[BusinessRule] = field(default_factory=list)
    features: list[Feature] = field(default_factory=list)
    appendix: list[str] = field(default_factory=list)
    signatures: list[Signature] = field(default_factory=list)
    present_sections: list[str] = field(default_factory=list)
    missing_sections: list[str] = field(default_factory=list)


def parse_product_spec(path: str | Path) -> ProductSpec:
    document = Document(str(path))
    spec = ProductSpec()
    seen: set[str] = set()
    section = ""
    subsection = ""
    feature: Feature | None = None

    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = _heading_level(block)
            if level == 1:
                section = _match_name(text, _SECTIONS, _SECTION_ALIASES)
                subsection = ""
                feature = None
                if section:
                    seen.add(section)
                continue
            if level == 2 and section:
                feature = None
                if section == "业务需求":
                    started = _start_feature(text)
                    if started is not None:
                        feature = started
                        spec.features.append(feature)
                        subsection = ""
                        continue
                subsection = _match_name(text, _SUBSECTIONS.get(section, ()))
                continue
            _consume_paragraph(spec, section, subsection, feature, text)
        elif isinstance(block, Table):
            _consume_table(spec, section, subsection, feature, block)

    spec.features = [item for item in spec.features if _keep_feature(item)]
    spec.revisions = [item for item in spec.revisions if not is_blank(item.number)]
    spec.definitions = [item for item in spec.definitions if not is_blank(item.term)]
    spec.references = [item for item in spec.references if not is_blank(item.name)]
    spec.files = [item for item in spec.files if not is_blank(item.path)]
    spec.users = [item for item in spec.users if not is_blank(item.name)]
    spec.modules = [item for item in spec.modules if not is_blank(item.module_id) or not is_blank(item.name)]
    spec.obvious = _unique(spec.obvious)
    spec.rules = [item for item in spec.rules if not is_blank(item.description) or not is_blank(item.sample_output)]
    spec.appendix = _unique(spec.appendix)
    spec.signatures = [item for item in spec.signatures if not is_blank(item.role)]
    spec.present_sections = [name for name in _SECTIONS if name in seen]
    spec.missing_sections = [name for name in _SECTIONS if name not in seen]
    return spec


def current_version(spec: ProductSpec) -> str:
    for revision in reversed(spec.revisions):
        if not is_blank(revision.version_after):
            return revision.version_after.strip()
    return ""


def split_items(text: str) -> list[str]:
    found = []
    for part in re.split(r"[；;\n]+", text or ""):
        item = part.strip()
        if is_blank(item) or item in found:
            continue
        found.append(item)
    return found


def render_formal(spec: ProductSpec) -> str:
    """给 coding agent 的正文。附录不在其中。"""
    lines = ["下面是已经确认的正式需求。只实现这些内容。", ""]
    _line(lines, "软件名称", spec.scope.software_name or spec.cover.module)
    _line(lines, "编写目的", spec.purpose)
    _line(lines, "要做的事", spec.scope.does)
    _line(lines, "应用目标", spec.scope.goal)
    _line(lines, "项目背景", spec.background)
    _line(lines, "项目周期", spec.schedule)
    _line(lines, "产品描述", spec.description)
    runtime = spec.runtime
    _line(lines, "exe 文件名", runtime.exe_name)
    _line(lines, "主窗口标题", runtime.window_title)
    _line(lines, "如何启动", runtime.how_to_start)
    _line(lines, "启动参数", runtime.launch_args)
    _line(lines, "如何退出", runtime.how_to_exit)
    _line(lines, "工作目录", runtime.workdir)
    _line(lines, "是否需要登录", runtime.needs_login)
    _line(lines, "测试账号", runtime.account)
    _line(lines, "分辨率", runtime.resolution)
    if spec.files:
        lines.append("文件：")
        for item in spec.files:
            lines.append(f"- {item.role} {item.path} {item.description}".strip())
    if spec.definitions:
        lines.append("用语：")
        for item in spec.definitions:
            ui_name = item.ui_name or item.term
            lines.append(f"- {item.term}：{item.meaning}。屏幕上写成「{ui_name}」。")
    if spec.users:
        lines.append("使用者：")
        for item in spec.users:
            lines.append(
                f"- {item.name}（{item.who}）要完成：{item.task}。在意：{item.cares}。不该做：{item.avoid}。"
            )
    if spec.modules:
        lines.append("功能模块：")
        for item in spec.modules:
            lines.append(f"- {item.module_id} {item.name}：{item.summary}。优先级 {item.priority}。")
    if spec.obvious:
        lines.append("不可接受的现象：")
        for item in spec.obvious:
            lines.append(f"- {item}")
    rules = [*spec.rules, *[rule for feature in spec.features for rule in feature.rules]]
    if rules:
        lines.append("业务规则：")
        for rule in rules:
            lines.append(f"- {rule.description}。输入样例：{rule.sample_input}。应显示：{rule.sample_output}。")
    excluded = split_items(spec.scope.does_not)
    for feature in spec.features:
        excluded.extend(item for item in split_items(feature.excluded) if item not in excluded)
    if excluded:
        lines.append("明确不做：")
        for item in excluded:
            lines.append(f"- {item}")
    for feature in spec.features:
        lines.append("")
        lines.append(f"{feature.feature_id} {feature.name}")
        _line(lines, "所属模块", feature.module)
        _line(lines, "功能介绍", feature.intro)
        _line(lines, "入口", feature.entry)
        _line(lines, "前置状态", feature.precondition)
        if feature.steps:
            lines.append("流程：")
            for step in feature.steps:
                lines.append(f"- （{step.applicability or '全部'}）{step.text}")
        if feature.items:
            lines.append("数据项：")
            for item in feature.items:
                lines.append(
                    f"- {item.label}（{item.kind}，必填 {item.required}，来源 {item.source}）。"
                    f"合法例子：{item.valid_example}。非法例子：{item.invalid_example}。边界：{item.boundary}。"
                    f"{item.note}"
                )
        if feature.screens:
            lines.append("界面上会看到：")
            for screen in feature.screens:
                lines.append(f"- （{screen.applicability or '合法'}）{screen.path}：{screen.sees}")
        _line(lines, "做得不到位时", feature.unacceptable)
    lines.append("")
    lines.append("附录没有写在上面。不要实现附录，也不要自行增加功能。")
    lines.append("请用 Python 和 PyQt 在当前目录实现这个桌面程序。只用标准控件。")
    lines.append("按钮、输入框、下拉框、表格和窗口标题使用上面写出的可见文字，不要改名，不要靠坐标定位。")
    return "\n".join(lines)


def build_product_template(path: str | Path) -> None:
    _build(path, example=False)


def build_product_example(path: str | Path) -> None:
    _build(path, example=True)


def _build(path: str | Path, example: bool) -> None:
    data = _example_data() if example else None
    document = Document()
    _set_song(document)
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    section.different_first_page_header_footer = True
    _header_line(section.header.paragraphs[0], "产品需求说明书", WD_ALIGN_PARAGRAPH.RIGHT)
    _header_line(section.footer.paragraphs[0], "附录不作为正式需求，也不作为验收标准。", WD_ALIGN_PARAGRAPH.CENTER)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("产品需求说明书" if data is None else data["cover"]["文档标题"])
    _font(run, 16, bold=True)
    document.add_paragraph("首页不设页眉页脚。从第 2 页起使用页眉和页脚。正文用宋体五号、1.5 倍行距。")
    document.add_paragraph("附录只保留背景，不作为正式需求，也不作为验收标准。确认、开发和测试都不采用附录。")

    document.add_heading("首页", level=1)
    cover = data["cover"] if data else {label: PLACEHOLDER for label, _ in _COVER_FIELDS}
    _kv(document, [(label, cover[label]) for label, _ in _COVER_FIELDS])

    document.add_heading("修订记录", level=1)
    revision_header = ("编号", "章节名称", "修订内容简述", "修订日期", "修订前版本号", "修订后版本号", "修订人", "批准人")
    revisions = data["revisions"] if data else [("1", PLACEHOLDER, PLACEHOLDER, PLACEHOLDER, "—", PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, revision_header, revisions)

    document.add_heading("引言", level=1)
    document.add_heading("编写目的", level=2)
    document.add_paragraph(data["purpose"] if data else PLACEHOLDER)
    document.add_heading("范围", level=2)
    scope = data["scope"] if data else {label: PLACEHOLDER for label, _ in _SCOPE_FIELDS}
    _kv(document, [(label, scope[label]) for label, _ in _SCOPE_FIELDS])
    document.add_heading("定义", level=2)
    definitions = data["definitions"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("术语", "含义", "界面叫法"), definitions)
    document.add_heading("参考资料", level=2)
    references = data["references"] if data else [(PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("资料名称", "说明"), references)

    document.add_heading("项目概述", level=1)
    document.add_heading("项目背景", level=2)
    document.add_paragraph(data["background"] if data else PLACEHOLDER)
    document.add_heading("项目周期", level=2)
    document.add_paragraph(data["schedule"] if data else PLACEHOLDER)
    document.add_heading("产品描述", level=2)
    document.add_paragraph(data["description"] if data else PLACEHOLDER)
    document.add_heading("运行方式", level=2)
    runtime = data["runtime"] if data else {label: PLACEHOLDER for label, _ in _RUNTIME_FIELDS}
    _kv(document, [(label, runtime[label]) for label, _ in _RUNTIME_FIELDS])
    files = data["files"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("文件角色", "路径", "说明"), files)
    document.add_heading("使用者", level=2)
    users = data["users"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("角色", "是谁", "要完成的事", "在意什么", "不该做的事"), users)
    document.add_heading("产品功能", level=2)
    modules = data["modules"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("模块编号", "模块名称", "功能简述", "优先级"), modules)

    document.add_heading("业务需求", level=1)
    document.add_heading("不可接受的现象", level=2)
    obvious = data["obvious"] if data else [(PLACEHOLDER,)]
    _grid(document, ("现象",), obvious)
    document.add_heading("业务规则", level=2)
    rules = data["rules"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("规则说明", "输入样例", "应显示的结果"), rules)
    features = data["features"] if data else [_blank_feature()]
    for feature in features:
        _add_feature(document, feature)

    document.add_page_break()
    document.add_heading("附录", level=1)
    document.add_paragraph(APPENDIX_NOTE)
    for line in data["appendix"] if data else ["把调研时听到、但这次不承诺的内容写在这里。"]:
        document.add_paragraph(line)

    document.add_heading("签字", level=1)
    signatures = data["signatures"] if data else [("产品经理", PLACEHOLDER, PLACEHOLDER), ("客户", PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("角色", "姓名", "日期"), signatures)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _example_data() -> dict:
    return {
        "cover": {
            "公司名称": "示例公司",
            "文档标题": "报销录入产品需求说明书",
            "文档编号": "PRD-BX-001",
            "编写人": "产品经理",
            "模块名称": "报销录入",
            "部门": "产品",
            "保密等级": "内部",
            "日期": "2026-10-08",
            "版权说明": "仅供确认、开发和测试使用",
        },
        "revisions": [("1", "全文", "首版", "2026-10-08", "—", "1.0.0", "产品经理", PLACEHOLDER)],
        "purpose": "验证报销明细的录入、合计与保存",
        "scope": {
            "软件名称": "报销录入",
            "软件要做的事": "录入差旅明细，计算合计，并保存到文件",
            "软件不做的事": "不测附件上传；不测同时打开两个程序；不测打印",
            "应用目标": "让报销人自己完成一张差旅报销单",
        },
        "definitions": [
            ("报销明细", "一次差旅里的一条花费", "明细表"),
            ("合计", "全部明细金额相加，保留两位小数", "合计"),
        ],
        "references": [("报销制度说明", "金额保留两位小数")],
        "background": "报销人现在用表格记差旅花费，合计容易算错。",
        "schedule": "客户确认之后开始开发。",
        "description": "一个桌面窗口。报销人添加明细，窗口底部显示合计，保存后写出文本文件。",
        "runtime": {
            "exe 文件名": "ExpenseEntry.exe",
            "主窗口标题": "报销录入",
            "如何启动": "双击 ExpenseEntry.exe",
            "启动参数": "无",
            "如何退出": "点击「退出」",
            "工作目录": "与 exe 相同",
            "是否需要登录": "否",
            "测试账号": "无",
            "分辨率": "1920×1080",
        },
        "files": [
            ("输入", "samples/cities.txt", "城市可选项包含「上海」"),
            ("输出", "输出/报销单.txt", "点击保存后生成，内含合计"),
        ],
        "users": [("普通员工", "报销人", "录入一条差旅明细并保存", "合计是否正确、能否保存", "不使用管理员菜单")],
        "modules": [("M-001", "报销单", "录入明细并保存", "高")],
        "obvious": [("合计为负数时，「保存」仍然可以点击",)],
        "rules": [("合计等于明细金额之和，保留两位小数", "128.50 与 71.50", "200.00")],
        "features": [
            {
                "id": "F-001",
                "name": "录入明细",
                "module": "M-001",
                "intro": "添加一条差旅明细",
                "entry": "主窗口",
                "precondition": "主窗口已打开，明细表为空",
                "steps": [("全部", "点击「添加」")],
                "items": [
                    ("事由", "文本", "手填", "差旅原因", "是", "出差车票", PLACEHOLDER, PLACEHOLDER),
                    ("金额", "数字", "手填", "大于 0", "是", "128.50", "-1", "0.01"),
                    ("城市", "下拉", "samples/cities.txt", "选项来自城市列表", "是", "上海", PLACEHOLDER, PLACEHOLDER),
                ],
                "screens": [
                    ("明细表", "出现一行，事由为「出差车票」，金额为「128.50」，城市为「上海」", "合法"),
                    ("提示", "请输入大于 0 的金额", "非法"),
                    ("明细表", "出现一行，金额列等于刚才输入的金额", "边界"),
                ],
                "boundary": PLACEHOLDER,
                "unacceptable": "请输入大于 0 的金额",
                "excluded": "不测附件上传",
            },
            {
                "id": "F-002",
                "name": "保存报销单",
                "module": "M-001",
                "intro": "把合计写入输出文件",
                "entry": "主窗口「保存」",
                "precondition": "主窗口已打开",
                "steps": [("合法", "点击「保存」"), ("非法", "在明细表为空时点击「保存」")],
                "items": [],
                "screens": [
                    ("输出文件", "文件「输出/报销单.txt」包含「合计：128.50」", "合法"),
                    ("提示", "请先添加明细", "非法"),
                ],
                "boundary": PLACEHOLDER,
                "unacceptable": "请先添加明细",
                "excluded": PLACEHOLDER,
            },
        ],
        "appendix": ["调研时客户提过要做打印，本次不纳入。"],
        "signatures": [("产品经理", PLACEHOLDER, PLACEHOLDER), ("客户", PLACEHOLDER, PLACEHOLDER)],
    }


def _blank_feature() -> dict:
    return {
        "id": "F-001",
        "name": PLACEHOLDER,
        "module": PLACEHOLDER,
        "intro": PLACEHOLDER,
        "entry": PLACEHOLDER,
        "precondition": PLACEHOLDER,
        "steps": [("全部", PLACEHOLDER)],
        "items": [(PLACEHOLDER, "文本", PLACEHOLDER, PLACEHOLDER, "是", PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)],
        "screens": [(PLACEHOLDER, PLACEHOLDER, "合法")],
        "boundary": PLACEHOLDER,
        "unacceptable": PLACEHOLDER,
        "excluded": PLACEHOLDER,
    }


def _add_feature(document, feature: dict) -> None:
    document.add_heading(f"{feature['id']} {feature['name']}", level=2)
    document.add_paragraph(f"所属模块：{feature['module']}")
    document.add_paragraph(f"功能介绍：{feature['intro']}")
    document.add_paragraph(f"入口：{feature['entry']}")
    document.add_paragraph(f"前置状态：{feature['precondition']}")
    document.add_paragraph("流程")
    _grid(document, ("序号", "适用", "步骤"), [(str(index), item[0], item[1]) for index, item in enumerate(feature["steps"], start=1)])
    document.add_paragraph("数据项")
    rows = feature["items"] or [(PLACEHOLDER, "文本", PLACEHOLDER, PLACEHOLDER, "是", PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _grid(document, ("字段", "字段类型", "数据来源", "备注", "必填", "合法例子", "非法例子", "边界"), rows)
    document.add_paragraph("界面")
    _grid(document, ("功能路径", "会看到什么", "适用"), feature["screens"])
    document.add_paragraph(f"边界时会看到：{feature['boundary']}")
    document.add_paragraph(f"不可接受：{feature['unacceptable']}")
    document.add_paragraph(f"本次不做：{feature['excluded']}")


def _consume_paragraph(spec: ProductSpec, section: str, subsection: str, feature: Feature | None, text: str) -> None:
    if section == "附录":
        spec.appendix.append(text)
        return
    if section == "引言" and subsection == "编写目的":
        spec.purpose = _join(spec.purpose, text)
        return
    if section == "项目概述" and subsection == "项目背景":
        spec.background = _join(spec.background, text)
        return
    if section == "项目概述" and subsection == "项目周期":
        spec.schedule = _join(spec.schedule, text)
        return
    if section == "项目概述" and subsection == "产品描述":
        spec.description = _join(spec.description, text)
        return
    if section == "业务需求" and subsection == "不可接受的现象" and _ITEM_PREFIX.match(text):
        spec.obvious.append(_clean_item(text))
        return
    if feature is None:
        return
    key, value = _split_label(text)
    attr = _FEATURE_LABELS.get(key)
    if not attr:
        return
    if attr == "intro":
        setattr(feature, attr, _join(getattr(feature, attr), value))
    else:
        setattr(feature, attr, value)


def _consume_table(spec: ProductSpec, section: str, subsection: str, feature: Feature | None, table: Table) -> None:
    rows = _table_rows(table)
    if len(rows) < 2:
        return
    header = [_compact(cell) for cell in rows[0]]
    data = rows[1:]
    if section == "首页" and header[:2] == ["项目", "内容"]:
        _apply(spec.cover, data, _COVER_FIELDS)
        return
    if section == "修订记录" and header[:4] == ["编号", "章节名称", "修订内容简述", "修订日期"]:
        for row in data:
            padded = _pad(row, 8)
            spec.revisions.append(Revision(*[cell.strip() for cell in padded]))
        return
    if section == "引言" and subsection == "范围" and header[:2] == ["项目", "内容"]:
        _apply(spec.scope, data, _SCOPE_FIELDS)
        return
    if header[:2] == ["术语", "含义"]:
        for row in data:
            padded = _pad(row, 3)
            spec.definitions.append(Definition(padded[0].strip(), padded[1].strip(), padded[2].strip()))
        return
    if header[:2] == ["资料名称", "说明"]:
        for row in data:
            padded = _pad(row, 2)
            spec.references.append(Reference(padded[0].strip(), padded[1].strip()))
        return
    if section == "项目概述" and subsection == "运行方式" and header[:2] == ["项目", "内容"]:
        _apply(spec.runtime, data, _RUNTIME_FIELDS)
        return
    if header[:3] == ["文件角色", "路径", "说明"]:
        for row in data:
            padded = _pad(row, 3)
            spec.files.append(SpecFile(*[cell.strip() for cell in padded]))
        return
    if header[:5] == ["角色", "是谁", "要完成的事", "在意什么", "不该做的事"]:
        for row in data:
            padded = _pad(row, 5)
            spec.users.append(UserRole(*[cell.strip() for cell in padded]))
        return
    if header[:4] == ["模块编号", "模块名称", "功能简述", "优先级"]:
        for row in data:
            padded = _pad(row, 4)
            spec.modules.append(Module(*[cell.strip() for cell in padded]))
        return
    if section == "业务需求" and subsection == "不可接受的现象" and header[:1] == ["现象"]:
        spec.obvious.extend(row[0].strip() for row in data if row and row[0].strip())
        return
    if header[:3] == ["规则说明", "输入样例", "应显示的结果"]:
        target = feature.rules if feature is not None else spec.rules
        for row in data:
            padded = _pad(row, 3)
            target.append(BusinessRule(*[cell.strip() for cell in padded]))
        return
    if feature is None:
        if section == "签字" and header[:3] == ["角色", "姓名", "日期"]:
            for row in data:
                padded = _pad(row, 3)
                spec.signatures.append(Signature(*[cell.strip() for cell in padded]))
        return
    if header[:3] == ["序号", "适用", "步骤"]:
        for row in data:
            padded = _pad(row, 3)
            feature.steps.append(FlowStep(padded[1].strip() or "全部", padded[2].strip()))
        return
    if header[:8] == ["字段", "字段类型", "数据来源", "备注", "必填", "合法例子", "非法例子", "边界"]:
        for row in data:
            padded = _pad(row, 8)
            feature.items.append(DataItem(*[cell.strip() for cell in padded]))
        return
    if header[:3] == ["功能路径", "会看到什么", "适用"]:
        for row in data:
            padded = _pad(row, 3)
            feature.screens.append(Screen(padded[0].strip(), padded[1].strip(), padded[2].strip() or "合法"))


def _apply(target, data, fields) -> None:
    mapping = {_compact(label): attr for label, attr in fields}
    for row in data:
        if len(row) < 2:
            continue
        attr = mapping.get(_compact(row[0]))
        if attr:
            setattr(target, attr, row[1].strip())


def _start_feature(text: str) -> Feature | None:
    match = re.match(r"^(F[-_]\d+)\s*(.*)$", text.strip(), re.IGNORECASE)
    if not match:
        return None
    return Feature(feature_id=match.group(1).upper().replace("_", "-"), name=match.group(2).strip())


def _keep_feature(feature: Feature) -> bool:
    feature.steps = [step for step in feature.steps if not is_blank(step.text)]
    feature.items = [item for item in feature.items if not is_blank(item.label)]
    feature.screens = [item for item in feature.screens if not is_blank(item.path) and not is_blank(item.sees)]
    feature.rules = [item for item in feature.rules if not is_blank(item.description) or not is_blank(item.sample_output)]
    if any(not is_blank(getattr(feature, name)) for name in ("name", "intro", "precondition", "unacceptable", "excluded", "boundary_sees")):
        return True
    return bool(feature.steps or feature.items or feature.screens or feature.rules)


def _match_name(text: str, names: tuple[str, ...], aliases: dict[str, str] | None = None) -> str:
    candidates = [text, re.sub(r"^[\d.、\s]+", "", text).strip()]
    for candidate in candidates:
        compact = _compact(candidate)
        if aliases and compact in {_compact(key): value for key, value in aliases.items()}:
            return {_compact(key): value for key, value in aliases.items()}[compact]
        for name in names:
            if compact == _compact(name):
                return name
    return ""


def _join(current: str, extra: str) -> str:
    extra = (extra or "").strip()
    if is_blank(extra):
        return current
    if not current:
        return extra
    return f"{current}\n{extra}"


def _unique(items: list[str]) -> list[str]:
    found = []
    for item in items:
        text = item.strip()
        if is_blank(text) or text in found:
            continue
        found.append(text)
    return found


def _line(lines: list[str], label: str, value: str) -> None:
    if not is_blank(value):
        lines.append(f"{label}：{value.strip()}")


def _kv(document, rows):
    _grid(document, ("项目", "内容"), rows)


def _grid(document, headers, rows):
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _cell(table.rows[0].cells[index], header, bold=True, fill=_HEADER_FILL, color=RGBColor(255, 255, 255))
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            _cell(table.rows[row_index + 1].cells[col_index], value)
    document.add_paragraph()


def _cell(cell, text: str, bold: bool = False, fill: str | None = None, color: RGBColor | None = None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = paragraph.add_run("" if text is None else str(text))
    _font(run, 10.5, bold=bold, color=color)
    if fill:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)


def _set_song(document) -> None:
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    _style_font(normal)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Title", 16)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        _style_font(style)


def _style_font(style) -> None:
    properties = style.element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:ascii"), "宋体")
    fonts.set(qn("w:hAnsi"), "宋体")
    fonts.set(qn("w:eastAsia"), "宋体")


def _font(run, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    if color is not None:
        run.font.color.rgb = color
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:ascii"), "宋体")
    fonts.set(qn("w:hAnsi"), "宋体")
    fonts.set(qn("w:eastAsia"), "宋体")


def _header_line(paragraph, text: str, alignment) -> None:
    paragraph.alignment = alignment
    run = paragraph.add_run(text)
    _font(run, 9)


def main() -> None:
    root = Path(__file__).resolve().parents[2]
    build_product_template(root / "templates" / "产品需求说明书模板.docx")
    build_product_example(root / "templates" / "产品需求说明书示例.docx")


if __name__ == "__main__":
    main()
