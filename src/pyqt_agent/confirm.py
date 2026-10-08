"""把技术需求书改写成用户能逐条勾选的确认单。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pyqt_agent.cases import expand_cases
from pyqt_agent.docx_parser import parse_requirements
from pyqt_agent.models import UNTESTED_MISSING, Requirements, TestCase, Untested, is_blank

_SITUATION = re.compile(r"（([^=]+)=(.+)）\s*$")


def write_confirmation(requirements_path: str | Path, out_path: str | Path) -> Path:
    requirements = parse_requirements(requirements_path)
    destination = Path(out_path)
    if destination.suffix.lower() != ".docx":
        destination.mkdir(parents=True, exist_ok=True)
        destination = destination / "用户确认.docx"
    else:
        destination.parent.mkdir(parents=True, exist_ok=True)
    _render(requirements, destination)
    return destination


def _render(requirements: Requirements, path: Path) -> None:
    cases, untested = expand_cases(requirements)
    document = Document()
    _set_font(document)
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    app_name = _show(requirements.cover.app_name, "这套软件")
    title = document.add_heading(f"请确认：这是不是你要的{app_name}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "请不要阅读技术需求书。下面每一条都是你自己会做的事，以及做完后能在屏幕上或文件里亲眼看到的结果。"
    )
    document.add_paragraph(
        "请对每一条勾选「是这样」「不是这样」或「我要补充」。补充时用你自己的话写，不用改写成技术说法。"
        "只有你勾成「是这样」的内容，才会交给编程。没写在这里的功能，这次不会做。"
    )
    if not is_blank(requirements.cover.purpose):
        document.add_paragraph(f"这次准备做的事：{requirements.cover.purpose.strip()}")

    _add_roles(document, requirements)
    _add_glossary(document, requirements)
    document.add_heading("请逐条确认", level=1)
    if not cases:
        document.add_paragraph("现在还没有可以让你确认的功能。先不要把需求交给编程。")
    for index, case in enumerate(cases, start=1):
        _add_case(document, requirements, case, index)
    _add_rules(document, requirements)
    _add_obvious(document, requirements)
    _add_lists(
        document,
        "这次先不做",
        "下面这些这次不会做。如果其中有你其实需要的，请写下来。",
        _skipped(untested, requirements),
    )
    _add_lists(
        document,
        "你已经知道、这次先不改的情况",
        "如果不能接受，请直接写下来。",
        requirements.known_issues,
    )
    unclear = _unclear(untested)
    if unclear:
        document.add_heading("这些还没说清楚", level=1)
        document.add_paragraph("下面这些还写不出你能看到的结果，先不要当成已经答应你。")
        for item in unclear:
            document.add_paragraph(item, style="List Bullet")
    document.add_heading("请签名", level=1)
    document.add_paragraph("我确认：上面我勾成「是这样」的内容，就是我要的。没写到的，这次不会做。")
    document.add_paragraph("姓名：________________    日期：________________")

    document.add_page_break()
    document.add_heading("给产品经理留档（打印给用户前删掉这一页）", level=1)
    document.add_paragraph(
        "用户只确认前面的勾选。全部关键条目标成「是这样」之后，把原来的技术需求书交给 coding agent，不要把这份确认单当作编程规格。"
        "用户写了「不是这样」或补充时，先改技术需求书，再重新生成确认单，让用户再看一遍改过的那几条。"
    )
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _cell(table.rows[0].cells[0], "用户看到的条目", bold=True)
    _cell(table.rows[0].cells[1], "需求书里的编号", bold=True)
    for index, case in enumerate(cases, start=1):
        row = table.add_row()
        _cell(row.cells[0], f"{index}. {_headline(requirements, case)}")
        _cell(row.cells[1], case.id)
    document.save(str(path))


def _add_roles(document, requirements: Requirements) -> None:
    if not requirements.roles:
        return
    document.add_heading("谁在用", level=1)
    for role in requirements.roles:
        who = role.who.strip() if not is_blank(role.who) else role.name
        document.add_paragraph(f"{who}要完成的事：{_show(role.task, '（还没写）')}")
        if not is_blank(role.avoid):
            document.add_paragraph(f"这个人不会去做：{role.avoid.strip()}")


def _add_glossary(document, requirements: Requirements) -> None:
    pairs = [
        term
        for term in requirements.glossary
        if not is_blank(term.business) and not is_blank(term.ui_name) and term.business.strip() != term.ui_name.strip()
    ]
    if not pairs:
        return
    document.add_heading("屏幕上的叫法", level=1)
    document.add_paragraph("你口头说的词，和屏幕上印出来的词可能不一样。请核对是不是同一个东西。")
    for term in pairs:
        note = f"，{term.note.strip()}" if not is_blank(term.note) else ""
        document.add_paragraph(f"你说的「{term.business.strip()}」，屏幕上写成「{term.ui_name.strip()}」{note}。", style="List Bullet")


def _add_case(document, requirements: Requirements, case: TestCase, index: int) -> None:
    document.add_heading(f"{index}. {_headline(requirements, case)}", level=2)
    if case.precondition:
        document.add_paragraph(f"开始之前：{case.precondition}")
    document.add_paragraph("你会这样做：")
    if case.steps:
        for step_index, step in enumerate(case.steps, start=1):
            document.add_paragraph(f"{step_index}. {step}")
    else:
        document.add_paragraph("具体怎么点，需求书里还没写。")
    document.add_paragraph("这时你会看到：")
    for expected in case.expected:
        document.add_paragraph(expected, style="List Bullet")
    _choice_table(document)


def _add_rules(document, requirements: Requirements) -> None:
    written = [rule for rule in requirements.rules if not is_blank(rule.sample_input) or not is_blank(rule.sample_output)]
    if not written:
        return
    document.add_heading("请核对这笔账", level=1)
    document.add_paragraph("不用看公式。请看这个例子算出来的数是不是你要的。")
    for rule in written:
        description = rule.description.strip() if not is_blank(rule.description) else "计算结果"
        document.add_paragraph(f"{description}。例如：{rule.sample_input.strip()}，你应该看到 {rule.sample_output.strip()}。")
        _choice_table(document)


def _add_obvious(document, requirements: Requirements) -> None:
    if not requirements.obvious_errors:
        return
    document.add_heading("出现这些情况，算做错了", level=1)
    for item in requirements.obvious_errors:
        document.add_paragraph(item, style="List Bullet")
    document.add_paragraph("如果上面某一条其实可以接受，请写下来。")
    _choice_table(document)


def _add_lists(document, heading: str, intro: str, items: list[str]) -> None:
    if not items:
        return
    document.add_heading(heading, level=1)
    document.add_paragraph(intro)
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _choice_table(document) -> None:
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    headers = ("是这样", "不是这样", "我要补充（用你自己的话写）")
    for index, header in enumerate(headers):
        _cell(table.rows[0].cells[index], header, bold=True)
    _cell(table.rows[1].cells[0], "☐")
    _cell(table.rows[1].cells[1], "☐")
    _cell(table.rows[1].cells[2], "")
    document.add_paragraph()


def _headline(requirements: Requirements, case: TestCase) -> str:
    name = case.feature_id
    for feature in requirements.features:
        if feature.feature_id == case.feature_id and not is_blank(feature.name):
            name = feature.name.strip()
            break
    situation = ""
    matched = _SITUATION.search(case.title)
    if matched:
        situation = f"{matched.group(1)}填成 {matched.group(2)}"
    if case.kind == "valid":
        return f"{name}：按平时的做法做一次"
    if case.kind == "invalid":
        return f"{name}：{situation or '这样做得不对'}时，应该被拦住"
    if case.kind == "boundary":
        return f"{name}：{situation or '遇到边缘情况'}时"
    return name


def _skipped(untested: list[Untested], requirements: Requirements) -> list[str]:
    items = [item.item for item in untested if item.reason.startswith("需求标明不测")]
    for item in requirements.out_of_scope:
        if item not in items:
            items.append(item)
    return items


def _unclear(untested: list[Untested]) -> list[str]:
    return [item.item for item in untested if item.reason == UNTESTED_MISSING]


def _show(value: str, fallback: str) -> str:
    text = (value or "").strip()
    return fallback if is_blank(text) else text


def _set_font(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    _east_asia(normal.element.get_or_add_rPr(), "微软雅黑")


def _cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    run.font.name = "Calibri"
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    if bold:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), "1F4E79")
        cell._tc.get_or_add_tcPr().append(shading)
        run.font.color.rgb = RGBColor(255, 255, 255)


def _east_asia(properties, font_name: str) -> None:
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), font_name)
