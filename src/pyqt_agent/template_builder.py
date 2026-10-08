"""生成 Word 需求书模板和一份填写完整的示例。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pyqt_agent.models import (
    COVER_FIELDS,
    EXPECTED_HEADERS,
    FILE_HEADERS,
    INDEX_HEADERS,
    INPUT_HEADERS,
    ITEM_HEADERS,
    PLACEHOLDER,
    ROLE_HEADERS,
    RULE_HEADERS,
    STARTUP_FIELDS,
    STEP_HEADERS,
    TERM_HEADERS,
    WINDOW_HEADERS,
)

_HEADER_FILL = "1F4E79"
_HINT = RGBColor(0x66, 0x66, 0x66)


def build_template(path: str | Path) -> None:
    _build(path, example=False)


def build_example(path: str | Path) -> None:
    _build(path, example=True)


def _build(path: str | Path, example: bool) -> None:
    document = Document()
    _set_font(document)
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        header = section.header.paragraphs[0]
        header.text = "PyQt 黑盒测试需求书"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        footer.text = "控件只写界面上的文字，不写坐标，不写代码里的对象名。"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = document.add_heading("PyQt 黑盒测试需求书", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "本模板供测试智能体阅读。智能体只认下面这些标题和表头。"
        "某一节空着或整节删掉时，对应测试记为「需求未写，未测」，不会自行编造预期。"
    )
    document.add_heading("填写说明", level=1)
    for line in _instructions():
        document.add_paragraph(line, style="List Bullet")
    document.add_page_break()

    data = _example_data() if example else None
    _add_cover(document, data)
    _add_startup(document, data)
    _add_environment(document, data)
    _add_terms(document, data)
    _add_roles(document, data)
    _add_index(document, data)
    _add_features(document, data)
    _add_windows(document, data)
    _add_rules(document, data)
    _add_obvious(document, data)
    _add_known(document, data)
    _add_focus(document, data)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _instructions() -> list[str]:
    return [
        "同一控件全篇只用一个可见名称，用界面上的文字，不用代码对象名，不写坐标。",
        "预期结果必须是人能在界面或输出文件里核对的内容。",
        "没有样例数字的计算，不要求智能体判断对错。公式写在「计算和数据规则」，要核对的结果写进对应功能的预期。",
        "文件对话框写清要选的文件名，不写「随便选一个」。若程序靠命令行打开文件，把参数写进「启动参数」。",
        "预期结果表的「适用」只能填：合法、非法、边界、全部。「全部」只用于三种情况都要看到的事实，例如窗口标题不变。",
        "多个非法值或边界值用中文分号分开。应当拒绝的边界写成「应拒绝：-1」，它按非法输入测试。",
        "操作步骤写点击和菜单；具体填什么写在输入项表。已经写了非法例子时，不必再写一套非法步骤。",
        "没有单独输入值的失败路径（例如空表直接保存）把步骤的适用填成「非法」，并写清非法预期。",
        "若有登录，把登录写成功能明细的第一条，启动成功标志填登录窗口的标题，测试账号只写专用账号。",
        "增加功能时，复制「功能明细」里的一整节（标题 2），编号不要重复。标题格式：F-001 功能名称。",
        "模拟使用只走「角色」一节写过的人。建议关注点写在最后一节，不写成通过或失败。",
        "更细的说明和示例见仓库 docs/需求书填写说明.md。示例文档是 templates/需求书示例.docx。",
    ]


def _add_cover(document, data):
    document.add_heading("封面", level=1)
    cover = data["cover"] if data else {field: PLACEHOLDER for field in COVER_FIELDS}
    _kv_table(document, [(field, cover[field]) for field in COVER_FIELDS])


def _add_startup(document, data):
    document.add_heading("启动与就绪", level=1)
    document.add_paragraph("智能体用命令行里的 exe 启动程序。这里说明工作目录、参数，以及看到什么标题算窗口已就绪。")
    startup = data["startup"] if data else {field: PLACEHOLDER for field in STARTUP_FIELDS}
    _kv_table(document, [(field, startup[field]) for field in STARTUP_FIELDS])


def _add_environment(document, data):
    document.add_heading("环境与测试数据", level=1)
    resolution = data["resolution"] if data else PLACEHOLDER
    _kv_table(document, [("分辨率", resolution)])
    document.add_paragraph("文件角色填「输入」「可改动」或「输出」。路径相对工作目录。")
    rows = data["files"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _table(document, FILE_HEADERS, rows)


def _add_terms(document, data):
    document.add_heading("术语", level=1)
    document.add_paragraph("业务用词和界面叫法保持一对一，避免同一项在后文换名字。")
    rows = data["terms"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _table(document, TERM_HEADERS, rows)


def _add_roles(document, data):
    document.add_heading("角色", level=1)
    document.add_paragraph("模拟使用只走这里写过的角色。每个角色写清要完成的一件事，以及不该碰的功能。")
    rows = data["roles"] if data else [(PLACEHOLDER,) * 5]
    _table(document, ROLE_HEADERS, rows)


def _add_index(document, data):
    document.add_heading("功能目录", level=1)
    document.add_paragraph("每条功能一个编号。入口写菜单或按钮的可见文字。优先级填高、中或低。")
    rows = data["index"] if data else [(PLACEHOLDER,) * 5]
    _table(document, INDEX_HEADERS, rows)


def _add_features(document, data):
    document.add_heading("功能明细", level=1)
    document.add_paragraph(
        "每个功能单独一节，标题使用标题 2，格式为「F-001 功能名称」。"
        "下面是一节的结构，增加功能时整节复制。"
    )
    features = data["features"] if data else [_blank_feature()]
    for feature in features:
        document.add_heading(f"{feature['id']} {feature['name']}", level=2)
        document.add_paragraph(f"编号：{feature['id']}")
        document.add_paragraph(f"名称：{feature['name']}")
        document.add_paragraph(f"前置状态：{feature['precondition']}")
        document.add_paragraph("操作步骤")
        _table(document, STEP_HEADERS, feature["steps"])
        document.add_paragraph("输入项")
        _table(document, INPUT_HEADERS, feature["inputs"])
        document.add_paragraph("预期结果")
        _table(document, EXPECTED_HEADERS, feature["expected"])
        document.add_paragraph(f"失败时的提示：{feature['failure']}")
        document.add_paragraph(f"不测的分支：{feature['excluded']}")


def _add_windows(document, data):
    document.add_heading("界面与窗口", level=1)
    document.add_paragraph("只列窗口、对话框和菜单的名称与用途，不画布局。名称用窗口标题或菜单路径。")
    rows = data["windows"] if data else [(PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)]
    _table(document, WINDOW_HEADERS, rows)


def _add_rules(document, data):
    document.add_heading("计算和数据规则", level=1)
    document.add_paragraph("写出公式、舍入和样例。样例结果若要判定对错，还要写进对应功能的预期结果。")
    rows = data["rules"] if data else [(PLACEHOLDER,) * 4]
    _table(document, RULE_HEADERS, rows)


def _add_obvious(document, data):
    document.add_heading("算作明显错误的现象", level=1)
    rows = data["obvious"] if data else [(PLACEHOLDER,)]
    _table(document, ITEM_HEADERS, rows)


def _add_known(document, data):
    document.add_heading("已知问题与明确不测", level=1)
    document.add_heading("已知问题", level=2)
    known = data["known"] if data else [(PLACEHOLDER,)]
    _table(document, ITEM_HEADERS, known)
    document.add_heading("明确不测", level=2)
    skipped = data["skipped"] if data else [(PLACEHOLDER,)]
    _table(document, ITEM_HEADERS, skipped)


def _add_focus(document, data):
    document.add_heading("建议时请关注", level=1)
    rows = data["focus"] if data else [(PLACEHOLDER,)]
    _table(document, ITEM_HEADERS, rows)


def _blank_feature() -> dict:
    return {
        "id": "F-001",
        "name": PLACEHOLDER,
        "precondition": PLACEHOLDER,
        "steps": [("1", "全部", PLACEHOLDER)],
        "inputs": [(PLACEHOLDER, "文本", "是", PLACEHOLDER, PLACEHOLDER, PLACEHOLDER)],
        "expected": [("合法", PLACEHOLDER, PLACEHOLDER)],
        "failure": PLACEHOLDER,
        "excluded": PLACEHOLDER,
    }


def _example_data() -> dict:
    return {
        "cover": {
            "应用名称": "报销录入",
            "版本": "1.0.0",
            "exe 文件名": "ExpenseEntry.exe",
            "主窗口标题": "报销录入",
            "编写人": "测试负责人",
            "日期": "2026-10-08",
            "本次测试目的": "验证报销明细的录入、合计与保存",
        },
        "startup": {
            "工作目录": "与 exe 相同",
            "启动方式": "双击 ExpenseEntry.exe",
            "启动参数": PLACEHOLDER,
            "需要先打开的文件": "无",
            "启动成功标志": "报销录入",
            "如何正常退出": "点击「退出」",
            "是否需要登录": "否",
            "测试账号": "无",
        },
        "resolution": "1920×1080",
        "files": [
            ("输入", "samples/cities.txt", "城市可选项包含「上海」"),
            ("输出", "输出/报销单.txt", "点击保存后生成，内含合计"),
        ],
        "terms": [
            ("报销明细", "明细表", "主窗口中的表格"),
            ("合计", "合计", "窗口底部的只读金额"),
        ],
        "roles": [
            ("普通员工", "报销人", "录入一条差旅明细并保存", "合计是否正确、能否保存", "不使用管理员菜单"),
        ],
        "index": [
            ("F-001", "录入明细", "主窗口", "高", "添加一条报销明细"),
            ("F-002", "保存报销单", "主窗口「保存」", "高", "把合计写入输出文件"),
        ],
        "features": [
            {
                "id": "F-001",
                "name": "录入明细",
                "precondition": "主窗口已打开，明细表为空",
                "steps": [("1", "全部", "点击「添加」")],
                "inputs": [
                    ("事由", "文本", "是", "出差车票", PLACEHOLDER, PLACEHOLDER),
                    ("金额", "数字", "是", "128.50", "-1", "0.01"),
                    ("城市", "下拉", "是", "上海", PLACEHOLDER, PLACEHOLDER),
                ],
                "expected": [
                    ("合法", "明细表", "出现一行，事由为「出差车票」，金额为「128.50」，城市为「上海」"),
                    ("非法", "提示", "请输入大于 0 的金额"),
                    ("边界", "明细表", "出现一行，金额列等于刚才输入的金额"),
                ],
                "failure": "请输入大于 0 的金额",
                "excluded": "不测附件上传",
            },
            {
                "id": "F-002",
                "name": "保存报销单",
                "precondition": "主窗口已打开",
                "steps": [
                    ("1", "合法", "点击「保存」"),
                    ("2", "非法", "在明细表为空时点击「保存」"),
                ],
                "inputs": [],
                "expected": [
                    ("合法", "输出文件", "文件「输出/报销单.txt」包含「合计：128.50」"),
                    ("非法", "提示", "请先添加明细"),
                ],
                "failure": "请先添加明细",
                "excluded": PLACEHOLDER,
            },
        ],
        "windows": [
            ("报销录入", "窗口", "录入明细并保存"),
            ("提示", "对话框", "金额不合法或尚未添加明细时说明原因"),
            ("文件 > 退出", "菜单", "退出程序"),
        ],
        "rules": [
            ("R-001", "合计等于明细金额之和，保留两位小数", "128.50 与 71.50", "200.00"),
        ],
        "obvious": [("合计为负数时，「保存」仍然可以点击",)],
        "known": [("窗口最小化后恢复，明细表的滚动位置回到顶部",)],
        "skipped": [("不测同时打开两个程序",), ("不测打印",)],
        "focus": [
            ("完成保存要点的步骤是否过多",),
            ("提示文案是否能看懂",),
            ("金额填错之后能否改回",),
        ],
    }


def _kv_table(document, rows: list[tuple[str, str]]):
    _table(document, ("项目", "内容"), rows)


def _table(document, headers, rows):
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _write_cell(table.rows[0].cells[index], header, bold=True, fill=_HEADER_FILL, color=RGBColor(255, 255, 255))
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            _write_cell(table.rows[row_index + 1].cells[col_index], value)
    document.add_paragraph()
    return table


def _write_cell(cell, text: str, bold: bool = False, fill: str | None = None, color: RGBColor | None = None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    run.font.name = "Calibri"
    _east_asia(run, "微软雅黑")
    if color is not None:
        run.font.color.rgb = color
    if text == PLACEHOLDER:
        run.font.color.rgb = _HINT
        run.italic = True
    if fill:
        _shade(cell, fill)


def _set_font(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    _east_asia_style(normal, "微软雅黑")
    for style_name in ("Heading 1", "Heading 2", "Title"):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        _east_asia_style(style, "微软雅黑")


def _east_asia(run, font_name: str):
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:eastAsia"), font_name)


def _east_asia_style(style, font_name: str):
    properties = style.element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), font_name)


def _shade(cell, color: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def main():
    root = Path(__file__).resolve().parents[2]
    build_template(root / "templates" / "需求书模板.docx")
    build_example(root / "templates" / "需求书示例.docx")


if __name__ == "__main__":
    main()
