"""把产品需求说明书改写成客户能逐条确认的说明。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from pyqt_agent.docx_parser import _compact, _pad, _table_rows
from pyqt_agent.models import is_blank
from pyqt_agent.product_spec import (
    APPENDIX_NOTE,
    DataItem,
    Feature,
    ProductSpec,
    parse_product_spec,
    split_items,
)

PENDING = "（待确认）"


@dataclass
class Decision:
    item_id: str
    content: str
    verdict: str
    note: str


@dataclass
class _Item:
    item_id: str
    title: str
    content: str


def write_customer_doc(requirements_path: str | Path, out_path: str | Path) -> Path:
    spec = parse_product_spec(requirements_path)
    destination = Path(out_path)
    if destination.suffix.lower() != ".docx":
        destination.mkdir(parents=True, exist_ok=True)
        destination = destination / "客户确认.docx"
    else:
        destination.parent.mkdir(parents=True, exist_ok=True)
    _render(spec, destination)
    return destination


def read_decisions(path: str | Path) -> list[Decision]:
    document = Document(str(path))
    found: list[Decision] = []
    for table in document.tables:
        rows = _table_rows(table)
        if not rows:
            continue
        header = [_compact(cell) for cell in rows[0]]
        if header[:4] != ["编号", "内容", "客户结论", "补充"]:
            continue
        for row in rows[1:]:
            padded = _pad(row, 4)
            if is_blank(padded[0]) and is_blank(padded[1]):
                continue
            found.append(Decision(padded[0].strip(), padded[1].strip(), padded[2].strip(), padded[3].strip()))
    return found


def _render(spec: ProductSpec, path: Path) -> None:
    items = _items(spec)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(12)
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    name = spec.scope.software_name or spec.cover.module or "这套软件"
    title = document.add_heading(f"请确认：这是不是你要的{name}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("请不要阅读产品需求说明书。下面只写你会做的事，以及做完后能在屏幕上或文件里看到的结果。")
    document.add_paragraph("每一条请在文末表格里填「是这样」「不是这样」或「我要补充」。补充用你自己的话写。")
    document.add_paragraph("只有全部标成「是这样」之后，才会开始做这个软件。表格里没有的功能，这次不会做。")
    if not is_blank(spec.purpose):
        document.add_paragraph(f"这次准备做的事：{spec.purpose.strip()}")
    _names(document, spec)
    if spec.users:
        document.add_heading("谁在用", level=1)
        for user in spec.users:
            who = user.who or user.name
            document.add_paragraph(f"{who}要办成的事：{user.task}。在意的是：{user.cares}。不该做的事：{user.avoid}。")

    document.add_heading("请逐条看", level=1)
    if not items:
        document.add_paragraph("现在还没有可以确认的内容。先不要开始开发。")
    for item in items:
        document.add_heading(f"{item.item_id} {item.title}", level=2)
        document.add_paragraph(item.content)
        document.add_paragraph("请在文末表格填写：是这样、不是这样，或我要补充。")

    document.add_heading("确认表", level=1)
    document.add_paragraph("请在这张表里填写结论。客户结论请原样填写「是这样」「不是这样」或「我要补充」。还没看的条目保持「（待确认）」。")
    table = document.add_table(rows=1 + max(1, len(items)), cols=4)
    table.style = "Table Grid"
    for index, header in enumerate(("编号", "内容", "客户结论", "补充")):
        table.rows[0].cells[index].text = header
    if not items:
        table.rows[1].cells[0].text = "（无）"
    for index, item in enumerate(items, start=1):
        table.rows[index].cells[0].text = item.item_id
        table.rows[index].cells[1].text = item.content
        table.rows[index].cells[2].text = PENDING
        table.rows[index].cells[3].text = ""

    aside = [line for line in spec.appendix if line.strip() != APPENDIX_NOTE and not is_blank(line)]
    if aside:
        document.add_heading("先不用确认的闲谈", level=1)
        document.add_paragraph("下面这些来自附录，只是当时听到的话。这次不做，也不用在确认表里勾选。")
        for line in aside:
            document.add_paragraph(line)

    document.add_page_break()
    document.add_heading("给产品经理留档", level=1)
    document.add_paragraph("打印给客户之前删掉这一页。不要把这份确认单当作编程规格。")
    document.add_paragraph("开发以产品需求说明书的正文为准。确认表里每一条客户结论都是「是这样」之后，才能把正文交给编程。")
    document.add_paragraph("客户写了「不是这样」或补充时，先改产品需求说明书，再重新生成这份确认单。重新生成会把结论恢复成「（待确认）」。")
    for feature in spec.features:
        document.add_paragraph(f"{feature.feature_id} {feature.name}")
    document.save(str(path))


def _items(spec: ProductSpec) -> list[_Item]:
    items: list[_Item] = []

    def add(title: str, content: str) -> None:
        text = " ".join(content.split())
        if is_blank(text) or any(item.content == text for item in items):
            return
        items.append(_Item(f"C-{len(items) + 1:03d}", title, text))

    if not is_blank(spec.scope.does):
        add("要做的事", f"这套软件要做的事：{spec.scope.does.strip()}。")
    for user in spec.users:
        who = user.who or user.name
        add("使用的人", f"使用的人是{who}。要办成的事：{user.task}。不该做的事：{user.avoid}。")
    for feature in spec.features:
        _feature_items(feature, add)
    for rule in [*spec.rules, *[rule for feature in spec.features for rule in feature.rules]]:
        if is_blank(rule.sample_output) and is_blank(rule.description):
            continue
        add(
            "核对一个例子",
            f"按这个例子核对：输入{rule.sample_input}，你会看到{rule.sample_output}。{rule.description}。",
        )
    for item in split_items(spec.scope.does_not):
        add("这次先不做", f"这次先不做：{item}。")
    for feature in spec.features:
        for item in split_items(feature.excluded):
            add("这次先不做", f"这次先不做：{item}。")
    return items


def _feature_items(feature: Feature, add) -> None:
    title = feature.name or feature.feature_id
    intro = feature.intro.strip() if not is_blank(feature.intro) else title
    actions = _actions(feature)
    done = _sees(feature, "合法")
    happy = f"{intro}。"
    if actions:
        happy += f"你会这样做：{actions}。"
    if done:
        happy += f"办成了你会看到：{done}。"
    add(title, happy)

    missed = _sees(feature, "非法")
    if not missed and not is_blank(feature.unacceptable):
        missed = feature.unacceptable.strip()
    if missed:
        situation = _situation(feature, "invalid")
        lead = f"{situation}，" if situation else ""
        add(f"{title}做得不到位", f"{lead}做得不到位你会看到：{missed}。")

    edge = _sees(feature, "边界")
    if not edge and not is_blank(feature.boundary_sees):
        edge = feature.boundary_sees.strip()
    if edge:
        situation = _situation(feature, "boundary")
        lead = f"{situation}，" if situation else ""
        add(f"{title}的一个临界情况", f"{lead}这时你会看到：{edge}。")


def _actions(feature: Feature) -> str:
    parts = []
    for item in feature.items:
        if is_blank(item.valid_example):
            continue
        parts.append(_fill(item, item.valid_example))
    for step in feature.steps:
        if step.applicability in {"", "全部", "合法"} and not is_blank(step.text):
            parts.append(step.text.strip())
    return "，".join(parts)


def _situation(feature: Feature, kind: str) -> str:
    if kind == "invalid":
        for item in feature.items:
            value = _first_value(item.invalid_example)
            if value:
                return f"如果把「{item.label}」填成「{value}」"
        for step in feature.steps:
            if step.applicability == "非法" and not is_blank(step.text):
                return step.text.strip()
    if kind == "boundary":
        for item in feature.items:
            value = _first_plain_boundary(item.boundary)
            if value:
                return f"如果把「{item.label}」填成「{value}」"
    return ""


def _sees(feature: Feature, applicability: str) -> str:
    found = []
    for screen in feature.screens:
        if (screen.applicability or "合法") != applicability or is_blank(screen.sees):
            continue
        text = f"在{screen.path}里，{screen.sees.strip()}"
        if text not in found:
            found.append(text)
    return "；".join(found)


def _fill(item: DataItem, value: str) -> str:
    if item.kind.strip() == "下拉":
        return f"在「{item.label}」里选择「{value.strip()}」"
    return f"在「{item.label}」里填写「{value.strip()}」"


def _first_value(text: str) -> str:
    for part in split_items(text):
        cleaned = part
        for marker in ("应拒绝：", "应拒绝:"):
            if cleaned.startswith(marker):
                cleaned = cleaned[len(marker):].strip()
        if not is_blank(cleaned):
            return cleaned
    return ""


def _first_plain_boundary(text: str) -> str:
    for part in split_items(text):
        if part.startswith("应拒绝"):
            continue
        if not is_blank(part):
            return part
    return ""


def _names(document, spec: ProductSpec) -> None:
    lines = []
    for item in spec.definitions:
        ui_name = item.ui_name.strip()
        if is_blank(ui_name) or ui_name == item.term.strip():
            continue
        lines.append(f"你说的「{item.term}」，屏幕上写成「{ui_name}」。")
    if not lines:
        return
    document.add_heading("屏幕上的叫法", level=1)
    for line in lines:
        document.add_paragraph(line)
