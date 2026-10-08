"""把产品需求说明书的正文写成现有格式的测试需求说明书。附录不进入测试。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from pyqt_agent.models import is_blank
from pyqt_agent.product_spec import Feature, ProductSpec, Screen, current_version, parse_product_spec, split_items
from pyqt_agent.template_builder import _kv_table, _set_font, _table


def write_test_spec(requirements_path: str | Path, out_path: str | Path) -> Path:
    spec = parse_product_spec(requirements_path)
    destination = Path(out_path)
    if destination.suffix.lower() != ".docx":
        destination.mkdir(parents=True, exist_ok=True)
        destination = destination / "测试需求说明书.docx"
    else:
        destination.parent.mkdir(parents=True, exist_ok=True)
    _render(spec, destination)
    return destination


def _render(spec: ProductSpec, path: Path) -> None:
    document = Document()
    _set_font(document)
    app_name = _text(spec.scope.software_name or spec.cover.module)
    document.add_heading("封面", level=1)
    _kv_table(
        document,
        [
            ("应用名称", app_name),
            ("版本", _text(current_version(spec))),
            ("exe 文件名", _text(spec.runtime.exe_name)),
            ("主窗口标题", _text(spec.runtime.window_title)),
            ("编写人", _text(spec.cover.author)),
            ("日期", _text(spec.cover.date)),
            ("本次测试目的", _text(spec.purpose or spec.scope.does)),
        ],
    )

    document.add_heading("启动与就绪", level=1)
    _kv_table(
        document,
        [
            ("工作目录", _text(spec.runtime.workdir)),
            ("启动方式", _text(spec.runtime.how_to_start)),
            ("启动参数", _text(spec.runtime.launch_args)),
            ("需要先打开的文件", "无"),
            ("启动成功标志", _text(spec.runtime.window_title)),
            ("如何正常退出", _text(spec.runtime.how_to_exit)),
            ("是否需要登录", _text(spec.runtime.needs_login)),
            ("测试账号", _text(spec.runtime.account)),
        ],
    )

    document.add_heading("环境与测试数据", level=1)
    _kv_table(document, [("分辨率", _text(spec.runtime.resolution))])
    files = [(item.role, item.path, item.description) for item in spec.files] or [("无", "无", "无")]
    _table(document, ("文件角色", "路径", "说明"), files)

    document.add_heading("术语", level=1)
    terms = [(item.term, item.ui_name or item.term, item.meaning) for item in spec.definitions] or [("无", "无", "无")]
    _table(document, ("业务用词", "界面叫法", "说明"), terms)

    document.add_heading("角色", level=1)
    roles = [(item.name, item.who, item.task, item.cares, item.avoid) for item in spec.users] or [("无", "无", "无", "无", "无")]
    _table(document, ("角色", "是谁", "要完成的事", "在意什么", "不该碰到的功能"), roles)

    document.add_heading("功能目录", level=1)
    index_rows = []
    priorities = {item.module_id: item.priority for item in spec.modules}
    priorities.update({item.name: item.priority for item in spec.modules})
    for feature in spec.features:
        index_rows.append(
            (
                feature.feature_id,
                feature.name,
                _text(feature.entry, "主窗口"),
                _text(priorities.get(feature.module, ""), "中"),
                _text(feature.intro),
            )
        )
    _table(document, ("编号", "名称", "入口", "优先级", "说明"), index_rows or [("无", "无", "无", "无", "无")])

    document.add_heading("功能明细", level=1)
    if not spec.features:
        document.add_paragraph("无")
    for feature in spec.features:
        _add_feature(document, feature)

    document.add_heading("界面与窗口", level=1)
    _table(document, ("名称", "类型", "用途"), _windows(spec) or [("无", "无", "无")])

    document.add_heading("计算和数据规则", level=1)
    _table(document, ("规则编号", "说明", "输入样例", "应显示的结果"), _rules(spec) or [("无", "无", "无", "无")])

    document.add_heading("算作明显错误的现象", level=1)
    _table(document, ("条目",), [(item,) for item in spec.obvious] or [("无",)])

    document.add_heading("已知问题与明确不测", level=1)
    document.add_heading("已知问题", level=2)
    _table(document, ("条目",), [("无",)])
    document.add_heading("明确不测", level=2)
    excluded = split_items(spec.scope.does_not) or ["无"]
    _table(document, ("条目",), [(item,) for item in excluded])

    document.add_heading("建议时请关注", level=1)
    focus = [item.cares for item in spec.users if not is_blank(item.cares)] or ["无"]
    _table(document, ("条目",), [(item,) for item in focus])
    document.save(str(path))


def _add_feature(document: Document, feature: Feature) -> None:
    document.add_heading(f"{feature.feature_id} {feature.name}", level=2)
    document.add_paragraph(f"编号：{feature.feature_id}")
    document.add_paragraph(f"名称：{feature.name}")
    document.add_paragraph(f"前置状态：{_text(feature.precondition)}")
    steps = [
        (str(index), step.applicability or "全部", step.text)
        for index, step in enumerate(feature.steps, start=1)
        if not is_blank(step.text)
    ]
    if steps:
        _table(document, ("序号", "适用", "步骤"), steps)
    inputs = [
        (item.label, item.kind, item.required, item.valid_example, item.invalid_example, item.boundary)
        for item in feature.items
        if not is_blank(item.label)
    ]
    if inputs:
        _table(document, ("界面标签", "类型", "必填", "合法例子", "非法例子", "边界"), inputs)
    expected = _expected(feature)
    if expected:
        _table(document, ("适用", "观察位置", "预期"), expected)
    document.add_paragraph(f"失败时的提示：{_text(feature.unacceptable)}")
    document.add_paragraph(f"不测的分支：{_text(feature.excluded)}")


def _expected(feature: Feature) -> list[tuple[str, str, str]]:
    rows = []
    for screen in feature.screens:
        if is_blank(screen.sees):
            continue
        rows.append((screen.applicability or "合法", screen.path, screen.sees))
    if not is_blank(feature.boundary_sees) and not any(row[0] == "边界" for row in rows):
        rows.append(("边界", "界面", feature.boundary_sees))
    return rows


def _windows(spec: ProductSpec) -> list[tuple[str, str, str]]:
    rows = []
    seen = set()
    if not is_blank(spec.runtime.window_title):
        rows.append((spec.runtime.window_title, "窗口", "主窗口"))
        seen.add(spec.runtime.window_title)
    for feature in spec.features:
        for screen in feature.screens:
            if screen.path in seen:
                continue
            seen.add(screen.path)
            rows.append((screen.path, _window_kind(screen), screen.sees))
    return rows


def _window_kind(screen: Screen) -> str:
    path = screen.path
    if "菜单" in path or ">" in path or "＞" in path:
        return "菜单"
    if "提示" in path or "对话框" in path:
        return "对话框"
    if "文件" in path:
        return "文件"
    return "窗口"


def _rules(spec: ProductSpec) -> list[tuple[str, str, str, str]]:
    source = [*spec.rules, *[rule for feature in spec.features for rule in feature.rules]]
    rows = []
    for rule in source:
        if is_blank(rule.description) and is_blank(rule.sample_output):
            continue
        rows.append((f"R-{len(rows) + 1:03d}", rule.description, rule.sample_input, rule.sample_output))
    return rows


def _text(value: str, fallback: str = "无") -> str:
    text = (value or "").strip()
    if is_blank(text):
        return fallback
    return text
