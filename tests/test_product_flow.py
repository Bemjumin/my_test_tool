import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING

from pyqt_agent.cases import expand_cases
from pyqt_agent.cli import main
from pyqt_agent.code_agent import develop
from pyqt_agent.docx_parser import parse_requirements
from pyqt_agent.models import UNTESTED_MISSING, SessionResult
from pyqt_agent.product_agent import read_decisions, write_customer_doc
from pyqt_agent.product_spec import APPENDIX_NOTE, build_product_example, build_product_template, parse_product_spec, render_formal
from pyqt_agent.test_agent import run_test_agent
from pyqt_agent.test_spec import write_test_spec

ROOT = Path(__file__).resolve().parents[1]


def _text(path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _set_verdicts(path, verdict: str, first: str | None = None) -> None:
    document = Document(path)
    for table in document.tables:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if "客户结论" not in header:
            continue
        index = header.index("客户结论")
        for row_index, row in enumerate(table.rows[1:]):
            row.cells[index].text = first if first is not None and row_index == 0 else verdict
    document.save(path)


def test_product_template_uses_song_and_skips_the_first_header(tmp_path):
    path = tmp_path / "模板.docx"
    build_product_template(path)
    document = Document(path)
    section = document.sections[0]
    normal = document.styles["Normal"]
    assert section.different_first_page_header_footer
    assert section.header.paragraphs[0].text == "产品需求说明书"
    assert section.first_page_header.paragraphs[0].text == ""
    assert normal.font.size.pt == 10.5
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.ONE_POINT_FIVE
    assert "宋体" in normal.element.xml
    spec = parse_product_spec(path)
    assert spec.features == []
    assert APPENDIX_NOTE in spec.appendix


def test_example_is_the_only_source_and_checked_in(tmp_path):
    built = tmp_path / "built.docx"
    build_product_example(built)
    assert _text(built) == _text(ROOT / "templates" / "产品需求说明书示例.docx")
    spec = parse_product_spec(ROOT / "templates" / "产品需求说明书示例.docx")
    assert spec.scope.software_name == "报销录入"
    assert spec.runtime.exe_name == "ExpenseEntry.exe"
    assert [item.feature_id for item in spec.features] == ["F-001", "F-002"]
    assert any("调研时客户提过要做打印" in line for line in spec.appendix)
    formal = render_formal(spec)
    assert "事由" in formal
    assert "调研时客户提过" not in formal
    assert not any("打印" in item.name for item in spec.features)


def test_customer_sheet_is_plain_language(tmp_path):
    source = tmp_path / "产品需求.docx"
    build_product_example(source)
    out = write_customer_doc(source, tmp_path / "客户确认.docx")
    text = _text(out)
    customer, _, archive = text.partition("给产品经理留档")
    assert "请不要阅读产品需求说明书" in customer
    assert "出差车票" in customer
    assert "请输入大于 0 的金额" in customer
    assert "0.01" in customer
    assert "200.00" in customer
    assert "你说的「报销明细」，屏幕上写成「明细表」" in customer
    assert "屏幕上写成「合计」" not in customer
    assert "调研时客户提过要做打印" in customer
    assert "是这样" in customer and "不是这样" in customer and "我要补充" in customer
    for jargon in ("字段类型", "数据库", "枚举", "合法输入", "非法输入", "边界值", "控件", "适用", "需求未写", "验收"):
        assert jargon not in customer
    decisions = read_decisions(out)
    assert decisions
    assert all(item.verdict == "（待确认）" for item in decisions)
    assert not any("调研" in item.content for item in decisions)
    assert any("这次先不做" in item.content and "打印" in item.content for item in decisions)
    assert "不要把这份确认单当作编程规格" in archive

    _set_verdicts(out, "是这样")
    again = write_customer_doc(source, out)
    assert all(item.verdict == "（待确认）" for item in read_decisions(again))


def test_generated_test_spec_round_trips_without_the_appendix(tmp_path):
    source = tmp_path / "产品需求.docx"
    build_product_example(source)
    spec_path = write_test_spec(source, tmp_path / "测试需求说明书.docx")
    text = _text(spec_path)
    assert "调研时客户提过" not in text
    requirements = parse_requirements(spec_path)
    assert requirements.missing_sections == []
    cases, untested = expand_cases(requirements)
    by_id = {case.id: case for case in cases}
    assert set(by_id) == {
        "TC-F-001-valid",
        "TC-F-001-invalid-金额",
        "TC-F-001-boundary-金额",
        "TC-F-002-valid",
        "TC-F-002-invalid",
    }
    assert by_id["TC-F-001-valid"].steps == [
        "在「事由」输入「出差车票」",
        "在「金额」输入「128.50」",
        "在「城市」选择「上海」",
        "点击「添加」",
    ]
    assert by_id["TC-F-002-invalid"].steps == ["在明细表为空时点击「保存」"]
    assert any("128.50 与 71.50" in rule.sample_input for rule in requirements.rules)
    assert not any("打印" in feature.name for feature in requirements.features)
    assert any("打印" in item.item for item in untested)
    assert any("附件" in item.item for item in untested)
    assert UNTESTED_MISSING not in [item.reason for item in untested]


def test_coding_starts_only_after_every_item_is_accepted(tmp_path):
    source = tmp_path / "产品需求.docx"
    build_product_example(source)
    confirmed = write_customer_doc(source, tmp_path / "客户确认.docx")
    called = []

    def runner(prompt, cwd):
        called.append((prompt, cwd))

    blocked = develop(source, confirmed, tmp_path / "app", runner=runner)
    assert blocked.exit_code == 2
    assert called == []

    _set_verdicts(confirmed, "是这样", first="不是这样")
    rejected = develop(source, confirmed, tmp_path / "app", runner=runner)
    assert rejected.exit_code == 2
    assert called == []

    _set_verdicts(confirmed, "是这样")
    done = develop(source, confirmed, tmp_path / "app", runner=runner)
    assert done.exit_code == 0
    assert len(called) == 1
    prompt, cwd = called[0]
    assert cwd == str(tmp_path / "app")
    assert "事由" in prompt
    assert "不测打印" in prompt
    assert "调研时客户提过" not in prompt


def test_test_agent_writes_the_spec_then_calls_the_session(tmp_path):
    source = tmp_path / "产品需求.docx"
    build_product_example(source)
    seen = {}

    def session(exe, requirements, out):
        seen["exe"] = exe
        seen["requirements"] = requirements
        seen["out"] = out
        return SessionResult(report_path=str(Path(out) / "report.md"), exit_code=0)

    result = run_test_agent(str(source), "ExpenseEntry.exe", str(tmp_path / "报告"), session=session)
    assert result.exit_code == 0
    assert seen["exe"] == "ExpenseEntry.exe"
    assert Path(seen["requirements"]).name == "测试需求说明书.docx"
    requirements = parse_requirements(seen["requirements"])
    cases, _ = expand_cases(requirements)
    assert "TC-F-001-boundary-金额" in {case.id for case in cases}


def test_commands_for_the_three_agents(tmp_path, capsys):
    source = tmp_path / "产品需求.docx"
    build_product_example(source)
    confirmed = tmp_path / "客户确认.docx"
    assert main(["product", "--requirements", str(source), "--out", str(confirmed)]) == 0
    assert "客户确认单：" in capsys.readouterr().out
    assert main(["code", "--requirements", str(source), "--confirmed", str(confirmed), "--out", str(tmp_path / "app")]) == 2
    assert "是这样" in capsys.readouterr().out

    spec_path = tmp_path / "测试需求说明书.docx"
    assert main(["test-spec", "--requirements", str(source), "--out", str(spec_path)]) == 0
    assert parse_requirements(spec_path).missing_sections == []

    if sys.platform == "win32":
        return
    code = main(
        [
            "test-agent",
            "--requirements",
            str(source),
            "--exe",
            "ExpenseEntry.exe",
            "--out",
            str(tmp_path / "报告"),
        ]
    )
    assert code == 2
    assert (tmp_path / "报告" / "测试需求说明书.docx").is_file()
    assert "Windows" in (tmp_path / "报告" / "report.md").read_text(encoding="utf-8")
