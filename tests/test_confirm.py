from docx import Document

from pyqt_agent.cli import main
from pyqt_agent.confirm import write_confirmation
from pyqt_agent.template_builder import build_example


def _text(path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_confirmation_is_written_for_the_customer(tmp_path):
    requirements = tmp_path / "req.docx"
    build_example(requirements)
    out = write_confirmation(requirements, tmp_path / "用户确认.docx")
    text = _text(out)
    customer, _, appendix = text.partition("给产品经理留档")
    assert "请不要阅读技术需求书" in customer
    assert "按平时的做法做一次" in customer
    assert "出差车票" in customer
    assert "请输入大于 0 的金额" in customer
    assert "金额填成 0.01" in customer
    assert "200.00" in customer
    assert "附件上传" in customer
    assert "你说的「报销明细」，屏幕上写成「明细表」" in customer
    assert "是这样" in customer
    assert "不是这样" in customer
    assert "我要补充" in customer
    for jargon in ("合法输入", "非法输入", "边界值", "控件树", "适用", "需求未写"):
        assert jargon not in customer
    assert "TC-F-001-valid" in appendix
    assert "不要把这份确认单当作编程规格" in appendix
    original = _text(requirements)
    assert "请不要阅读技术需求书" not in original


def test_confirm_command_writes_the_file(tmp_path, capsys):
    requirements = tmp_path / "req.docx"
    build_example(requirements)
    code = main(["confirm", "--requirements", str(requirements), "--out", str(tmp_path / "out")])
    assert code == 0
    assert (tmp_path / "out" / "用户确认.docx").is_file()
    assert "用户确认单：" in capsys.readouterr().out
