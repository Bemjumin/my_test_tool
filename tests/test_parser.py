from pathlib import Path

from docx import Document

from pyqt_agent.docx_parser import parse_requirements
from pyqt_agent.template_builder import build_example, build_template

ROOT = Path(__file__).resolve().parents[1]


def test_example_document_has_the_expense_story():
    requirements = parse_requirements(ROOT / "templates" / "需求书示例.docx")
    assert requirements.missing_sections == []
    assert requirements.cover.app_name == "报销录入"
    assert requirements.cover.exe_name == "ExpenseEntry.exe"
    assert requirements.startup.ready_title == "报销录入"
    assert requirements.startup.how_to_exit == "点击「退出」"
    assert requirements.environment.resolution == "1920×1080"
    assert [item.path for item in requirements.environment.files] == ["samples/cities.txt", "输出/报销单.txt"]
    assert requirements.glossary[0].ui_name == "明细表"
    assert requirements.roles[0].name == "普通员工"
    assert [item.feature_id for item in requirements.feature_index] == ["F-001", "F-002"]
    assert [item.feature_id for item in requirements.features] == ["F-001", "F-002"]
    first = requirements.features[0]
    assert first.precondition == "主窗口已打开，明细表为空"
    assert first.inputs[1].invalid_example == "-1"
    assert first.inputs[1].boundary == "0.01"
    assert first.expected[0].applicability == "合法"
    assert "128.50" in first.expected[0].expected
    assert first.excluded_branches == "不测附件上传"
    assert requirements.rules[0].sample_output == "200.00"
    assert requirements.obvious_errors == ["合计为负数时，「保存」仍然可以点击"]
    assert len(requirements.known_issues) == 1
    assert len(requirements.out_of_scope) == 2
    assert len(requirements.suggestion_focus) == 3


def test_blank_template_has_every_section_and_no_real_feature():
    requirements = parse_requirements(ROOT / "templates" / "需求书模板.docx")
    assert requirements.missing_sections == []
    assert requirements.features == []
    assert requirements.roles == []
    assert requirements.obvious_errors == []
    assert "填写说明" not in requirements.present_sections


def test_repo_templates_match_the_builder(tmp_path):
    template = tmp_path / "template.docx"
    example = tmp_path / "example.docx"
    build_template(template)
    build_example(example)
    assert parse_requirements(template) == parse_requirements(ROOT / "templates" / "需求书模板.docx")
    assert parse_requirements(example) == parse_requirements(ROOT / "templates" / "需求书示例.docx")


def test_missing_section_is_reported(tmp_path):
    path = tmp_path / "partial.docx"
    document = Document()
    document.add_heading("封面", level=1)
    document.add_paragraph("应用名称：只有封面")
    document.save(path)
    requirements = parse_requirements(path)
    assert "封面" in requirements.present_sections
    assert "角色" in requirements.missing_sections
    assert "功能明细" in requirements.missing_sections
    assert requirements.cover.app_name == "只有封面"
